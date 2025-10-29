from __future__ import annotations
from typing import List, Optional
from pathlib import Path
import uuid
import textwrap
import logging

from pptx import Presentation
from PIL import Image, ImageDraw, ImageFont
from moviepy.editor import ImageClip, concatenate_videoclips, AudioFileClip

log = logging.getLogger(__name__)

DEFAULT_W = 1280
DEFAULT_H = 720


def extract_slide_texts(pptx_path: str) -> List[str]:
    prs = Presentation(pptx_path)
    slides_texts: List[str] = []
    for slide in prs.slides:
        parts: List[str] = []
        for shape in slide.shapes:
            try:
                if hasattr(shape, "text"):
                    t = shape.text.strip()
                    if t:
                        parts.append(t)
            except Exception:
                continue
        txt = "\n\n".join(parts)
        slides_texts.append(txt[:5000])
    return slides_texts


def _make_image_from_text(text: str, out_path: Path, width: int = DEFAULT_W, height: int = DEFAULT_H,
                          bg_color: str = "white", title: Optional[str] = None) -> str:
    img = Image.new("RGB", (width, height), color=bg_color)
    draw = ImageDraw.Draw(img)

    try:
        font = ImageFont.truetype("arial.ttf", 36)
        small = ImageFont.truetype("arial.ttf", 24)
    except Exception:
        font = ImageFont.load_default()
        small = ImageFont.load_default()

    y = 40
    if title:
        draw.text((60, y), title, fill="black", font=font)
        y += 70

    margin = 60
    lines = textwrap.wrap(text or "", width=80)
    for line in lines:
        draw.text((margin, y), line, fill="black", font=small)
        y += 28
        if y > height - 60:
            break

    out_path.parent.mkdir(parents=True, exist_ok=True)
    img.save(out_path, format="PNG")
    return str(out_path)


def render_slide_images(pptx_path: str, out_dir: str, prefix: str = "slide", width: int = DEFAULT_W,
                        height: int = DEFAULT_H) -> List[str]:
    out_dir_p = Path(out_dir)
    out_dir_p.mkdir(parents=True, exist_ok=True)
    slides = extract_slide_texts(pptx_path)
    paths: List[str] = []
    for i, txt in enumerate(slides, start=1):
        fname = out_dir_p / f"{prefix}_{i:03d}.png"
        title = None
        if '\n' in txt:
            title = txt.split('\n', 1)[0][:120]
        elif txt:
            title = txt[:120]
        _make_image_from_text(txt, fname, width=width, height=height, title=title)
        paths.append(str(fname))
    return paths


def create_video_from_images(image_paths: List[str], out_path: str, duration_per_slide: float = 4.0,
                             fps: int = 24, audio_path: Optional[str] = None) -> str:
    clips = []
    for img in image_paths:
        clip = ImageClip(img).set_duration(duration_per_slide)
        clips.append(clip)
    video = concatenate_videoclips(clips, method="compose")
    if audio_path and Path(audio_path).exists():
        try:
            audio = AudioFileClip(audio_path)
            video = video.set_audio(audio)
        except Exception as e:
            log.warning("Audio konnte nicht hinzugefügt werden: %s", e)

    out_p = Path(out_path)
    out_p.parent.mkdir(parents=True, exist_ok=True)
    video.write_videofile(str(out_p), fps=fps, codec="libx264", audio_codec="aac")
    return str(out_p)


def generate_video_from_pptx(pptx_path: str, out_dir: str, duration_per_slide: float = 4.0,
                             width: int = DEFAULT_W, height: int = DEFAULT_H) -> str:
    slides = extract_slide_texts(pptx_path)
    images_dir = Path(out_dir) / f"slides_{uuid.uuid4().hex[:8]}"
    images = render_slide_images(pptx_path, str(images_dir), width=width, height=height)
    out_mp4 = Path(out_dir) / (Path(pptx_path).stem + "_local_video.mp4")
    mp4 = create_video_from_images(images, str(out_mp4), duration_per_slide=duration_per_slide)
    return mp4
