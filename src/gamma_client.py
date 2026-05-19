"""Gamma-Präsentations-Client: generiert PPTX via Gamma-API oder lokalem python-pptx-Fallback."""
from __future__ import annotations
from typing import List, Dict, Any, Optional
import os
import json
import logging
from pathlib import Path

from .config import GAMMA_API_KEY, GAMMA_API_URL

log = logging.getLogger(__name__)

# Optional httpx for Gamma API calls
try:
    import httpx
except Exception:
    httpx = None  # type: ignore

# Local PPTX fallback
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except Exception:
    Presentation = None  # type: ignore

import requests


def _call_gamma_api(title: str, slides: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Sendet Titel und Folien an die Gamma-API und gibt die JSON-Antwort zurück."""
    if not GAMMA_API_URL or not GAMMA_API_KEY:
        raise RuntimeError("Gamma API nicht konfiguriert (GAMMA_API_URL/GAMMA_API_KEY fehlen)")
    if httpx is None:
        raise RuntimeError("httpx ist nicht installiert; Gamma-API-Aufrufe benötigen httpx")

    url = f"{GAMMA_API_URL.rstrip('/')}/presentations"
    payload = {"title": title, "slides": slides}
    headers = {"Authorization": f"Bearer {GAMMA_API_KEY}", "Content-Type": "application/json"}

    try:
        with httpx.Client(timeout=30.0) as client:
            r = client.post(url, headers=headers, json=payload)
            r.raise_for_status()
            return r.json()
    except Exception as e:
        log.exception("Gamma-API-Aufruf fehlgeschlagen: %s", e)
        raise


def _create_local_pptx(
    title: str,
    slides: List[Dict[str, Any]],
    out_dir: str = ".",
    template_path: Optional[str] = None,
    title_font_size: int = 32,
    body_font_size: int = 18,
    bullet_font_size: int = 14
) -> str:
    """Erstellt eine PPTX-Datei lokal via python-pptx als Gamma-Fallback und gibt den Pfad zurück."""
    if Presentation is None:
        raise RuntimeError("python-pptx nicht installiert; Installation in requirements.txt fehlt oder nicht installiert")

    # ensure output directory early (we may download images into it)
    out_dir_path = Path(out_dir or Path.cwd())
    out_dir_path.mkdir(parents=True, exist_ok=True)

    # Use provided template if available
    if template_path:
        try:
            tpl = Path(template_path)
            if tpl.exists():
                prs = Presentation(str(tpl))
            else:
                prs = Presentation()
        except Exception:
            prs = Presentation()
    else:
        prs = Presentation()

    # Title slide
    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = title
    subtitle = slide.placeholders[1]
    subtitle.text = "Erstellt mit Gamma-Fallback"

    # Content slides
    for s in slides:
        layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
        sl = prs.slides.add_slide(layout)
        try:
            sl.shapes.title.text = s.get("title", "")
            # Apply custom title font size
            if hasattr(sl.shapes.title, 'text_frame'):
                for paragraph in sl.shapes.title.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(title_font_size)
        except Exception:
            pass
        body_shape = None
        # find a body placeholder
        for ph in sl.placeholders:
            if ph.placeholder_format.type.name == "BODY":
                body_shape = ph
                break
        if body_shape is None:
            # fallback: try the second placeholder
            if len(sl.placeholders) >= 2:
                body_shape = sl.placeholders[1]

        content = s.get("content", "")
        if isinstance(content, list):
            lines = content
        else:
            # split into lines heuristically
            lines = [ln.strip() for ln in str(content).splitlines() if ln.strip()]
            if not lines:
                # fallback: split by sentence
                lines = [p.strip() for p in str(content).split(".") if p.strip()]

        # If the slide contains image markers like 'Bild: <url>', attempt to download
        # and insert the image; otherwise render text into the body placeholder.
        if body_shape is not None:
            tf = body_shape.text_frame
            tf.clear()
            for i, ln in enumerate(lines):
                # detect image marker
                if isinstance(ln, str) and ln.startswith("Bild:"):
                    img_url = ln.split("Bild:", 1)[1].strip()
                    if img_url:
                        try:
                            img_dir = out_dir_path / "images"
                            img_dir.mkdir(parents=True, exist_ok=True)
                            img_name = os.path.basename(img_url.split("?")[0]) or "image.jpg"
                            img_path = img_dir / img_name
                            if not img_path.exists():
                                # support local files (file:// or direct path)
                                local_src = None
                                if img_url.startswith("file://"):
                                    local_src = Path(img_url[len("file://"):])
                                else:
                                    pcheck = Path(img_url)
                                    if pcheck.exists():
                                        local_src = pcheck

                                if local_src is not None and local_src.exists():
                                    # copy local file into images dir
                                    from shutil import copyfile
                                    copyfile(str(local_src), str(img_path))
                                else:
                                    # download image
                                    r = requests.get(img_url, stream=True, timeout=30)
                                    r.raise_for_status()
                                    with open(img_path, "wb") as fh:
                                        for chunk in r.iter_content(chunk_size=1 << 14):
                                            if chunk:
                                                fh.write(chunk)
                            # add image to slide
                            left = Inches(1)
                            top = Inches(1.5)
                            width = Inches(7.5)
                            try:
                                sl.shapes.add_picture(str(img_path), left, top, width=width)
                            except Exception:
                                # fallback: put URL as text
                                p = tf.add_paragraph()
                                p.text = img_url
                                p.level = 1
                                p.font.size = Pt(bullet_font_size)
                            # continue to next line
                            continue
                        except Exception:
                            # if download or insert fails, fall back to placing URL as text
                            pass

                # regular text handling
                if i == 0:
                    p = tf.paragraphs[0]
                    p.text = ln
                    p.font.size = Pt(body_font_size)
                else:
                    p = tf.add_paragraph()
                    p.text = ln
                    p.level = 1
                    p.font.size = Pt(bullet_font_size)

    # ensure output directory
    out_dir_path = Path(out_dir or Path.cwd())
    out_dir_path.mkdir(parents=True, exist_ok=True)

    # create a safe filename
    safe_name = f"{title.replace(' ', '_')[:50]}.pptx"
    base = out_dir_path / safe_name
    i = 0
    dest = base
    while dest.exists():
        i += 1
        dest = base.with_name(f"{base.stem}_{i}.pptx")

    prs.save(str(dest))
    return str(dest)


def generate_presentation(
    title: str,
    answer_text: str,
    supports: List[Dict[str, Any]],
    use_gamma: bool = True,
    out_dir: str = "outputs/gamma",
    template_path: Optional[str] = None,
    title_font_size: int = 32,
    body_font_size: int = 18,
    bullet_font_size: int = 14
) -> Dict[str, Any]:
    """Erzeugt eine Präsentation aus Antworttext und Supports via Gamma-API oder lokalem PPTX-Fallback."""
    # 1) Erzeuge strukturierte Slides
    slides: List[Dict[str, Any]] = []

    # Titel-Slide (als Metadaten)
    slides.append({"title": title, "content": [ln for ln in answer_text.splitlines()[:3] if ln.strip()]})

    # Summary / Key points: heuristisch aus answer_text
    # simple: split into sentences and pick first 4
    sentences = [s.strip() for s in answer_text.replace('\n', ' ').split('. ') if s.strip()]
    key_points = sentences[:4]
    if key_points:
        slides.append({"title": "Kernaussagen", "content": key_points})

    # Add supports: each support -> one slide (limit)
    max_supports = 8
    for s in supports[:max_supports]:
        typ = s.get("type")
        if typ == "paragraph":
            title_s = f"Beleg: {s.get('paper_title','')} (S. {s.get('page')})"
            content = [s.get("text","")[:800]]
        elif typ == "figure":
            title_s = f"Abbildung: {s.get('paper_title','')} (S. {s.get('page')})"
            caption = s.get("caption","") or ""
            content = [caption]
            if s.get("image_uri"):
                content.append(f"Bild: {s.get('image_uri')}")
        else:
            title_s = s.get("paper_title","Quelle")
            content = [str(s.get("url") or s.get("doi") or "")] 
        slides.append({"title": title_s, "content": content})

    # 2) Versuche Gamma API
    if use_gamma and GAMMA_API_KEY and GAMMA_API_URL:
        try:
            api_resp = _call_gamma_api(title, slides)
            return {"method": "gamma", "result": api_resp, "slides": slides}
        except Exception as e:
            log.warning("Gamma API fehlgeschlagen, falle auf lokal PPTX zurück: %s", e)

    # 3) Local PPTX
    out_path = _create_local_pptx(
        title, 
        slides, 
        out_dir=(out_dir or Path.cwd()), 
        template_path=template_path,
        title_font_size=title_font_size,
        body_font_size=body_font_size,
        bullet_font_size=bullet_font_size
    )
    return {"method": "local", "result": {"path": out_path}, "slides": slides}

